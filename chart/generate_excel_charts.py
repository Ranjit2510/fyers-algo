import pandas as pd
import mplfinance as mpf
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from datetime import timedelta
import os

# ============================================================
# SETTINGS
# ============================================================

CSV_FILE = "data/banknifty_5min_data.csv"

CHART_FOLDER = "output-charts"

EXCEL_FILE = "output-charts/banknifty_weekday_charts.xlsx"

# Create chart folder
os.makedirs(CHART_FOLDER, exist_ok=True)

# ============================================================
# LOAD CSV
# ============================================================

df = pd.read_csv(CSV_FILE)

# Convert timestamp
df['Timestamp'] = pd.to_datetime(df['Timestamp'])

# Extract date
df['Date'] = df['Timestamp'].dt.date

# Set index
df.set_index('Timestamp', inplace=True)

# ============================================================
# CREATE EXCEL
# ============================================================

wb = Workbook()

ws = wb.active

ws.title = "BankNifty Charts"

# ============================================================
# WEEKDAY COLUMNS
# ============================================================

weekday_columns = {
    0: "A",   # Monday
    1: "D",   # Tuesday
    2: "G",   # Wednesday
    3: "J",   # Thursday
    4: "M"    # Friday
}

weekday_names = {
    0: "MONDAY",
    1: "TUESDAY",
    2: "WEDNESDAY",
    3: "THURSDAY",
    4: "FRIDAY"
}

# ============================================================
# HEADERS
# ============================================================

for day_num, col in weekday_columns.items():

    ws[f"{col}1"] = weekday_names[day_num]

# ============================================================
# TRACK ROW POSITION
# ============================================================

row_tracker = {
    0: 3,
    1: 3,
    2: 3,
    3: 3,
    4: 3
}

# ============================================================
# GET ALL DATES
# ============================================================

all_dates = sorted(df['Date'].unique())

# ============================================================
# PROCESS EACH DAY
# ============================================================

for current_date in all_dates:

    # Day dataframe
    day_df = df[df['Date'] == current_date]

    # Skip weekends
    weekday = pd.Timestamp(current_date).weekday()

    if weekday > 4:
        continue

    # ========================================================
    # CREATE CANDLE CHART
    # ========================================================

    chart_file = f"{CHART_FOLDER}/{current_date}.png"

    title = f"BankNifty - {current_date}"

    mpf.plot(
        day_df,
        type='candle',
        style='charles',
        title=title,
        volume=False,
        figsize=(10, 5),
        savefig=chart_file
    )

    # ========================================================
    # INSERT INTO EXCEL
    # ========================================================

    col = weekday_columns[weekday]

    row = row_tracker[weekday]

    # Add date text
    ws[f"{col}{row}"] = str(current_date)

    # Add image
    img = Image(chart_file)

    img.width = 600
    img.height = 300

    ws.add_image(img, f"{col}{row + 1}")

    # Move next row
    row_tracker[weekday] += 20

# ============================================================
# ADD HOLIDAYS
# ============================================================

start_date = min(all_dates)

end_date = max(all_dates)

all_calendar_dates = pd.date_range(start_date, end_date)

existing_dates = set(all_dates)

for date in all_calendar_dates:

    weekday = date.weekday()

    # Only weekdays
    if weekday > 4:
        continue

    if date.date() not in existing_dates:

        col = weekday_columns[weekday]

        row = row_tracker[weekday]

        ws[f"{col}{row}"] = f"{date.date()} - HOLIDAY"

        row_tracker[weekday] += 3

# ============================================================
# SAVE EXCEL
# ============================================================

wb.save(EXCEL_FILE)

print("\nExcel file generated successfully!")

print(f"\nSaved: {EXCEL_FILE}")